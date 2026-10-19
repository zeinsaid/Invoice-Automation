"""Write one PowerSoft import Excel per invoice (D-008), with flags as cell colours and comments (D-011)."""

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from .catalogue import Catalogue, Item
from .files import save_atomic
from .model import ORANGE, RED, SEVERITY, YELLOW, Flag

# The 20-column import format PowerSoft set up for us (D-012), exactly as in MASTER_ITEM_LIST.xlsx.
HEADER = ["Type", "Item Code", "Line No.", "Barcode No.", "Description", "Quantity", "Unit of Measure Code",
          "Unit Price Excl. VAT (Cost)", "Line Amount Excl. VAT", "Line Discount %", "Last Unit Price",
          "Retail price (incl. VAT)", "Category Code", "Category Name", "Department Code", "Department Name",
          "Supplier Code", "Supplier Name", "Brand", "VAT % (0,5,19)"]
WIDTHS = [6, 6, 9, 16, 36, 9, 6, 7, 6, 5, 6, 12, 8, 17, 7, 11, 10, 31, 16, 6]
NUMBER_FORMATS = {"D": "@", "F": "0.00", "H": "0.00", "I": "0.00", "J": "0.00", "K": "0.00", "L": "0.00"}

FILLS = {
    YELLOW: PatternFill("solid", fgColor="FFFF00"),
    ORANGE: PatternFill("solid", fgColor="FFC000"),
    RED: PatternFill("solid", fgColor="FF7C80"),
}
# Which cells a flag colours.
FLAG_COLUMNS = {"item": ["B", "E"], "quantity": ["F"], "cost": ["H"], "supplier": ["Q", "R"]}  # anything else: the whole row


@dataclass
class OutputRow:
    item: Item
    quantity: float  # selling pieces
    cost: float  # per selling piece, excl. VAT
    flags: list[Flag] = field(default_factory=list)
    invoice_lines: list[str] = field(default_factory=list)


def write_invoice_excel(path: Path, rows: list[OutputRow], catalogue: Catalogue, supplier: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="C0C0C0")

    for r, row in enumerate(rows, start=2):
        item = row.item
        code = int(item.code) if item.code.isdigit() else item.code
        # A new item gets this supplier; an existing item keeps PowerSoft's details (the tool never edits them).
        item_supplier = supplier if item.local else item.supplier
        ws.append([
            "Item", code, code, item.barcode, item.description, round(row.quantity, 3), "PCS",
            round(row.cost, 3), f"=F{r}*H{r}", 0, f"=H{r}", item.price,
            catalogue.category_codes.get(item.category), item.category or None, None, None,
            catalogue.supplier_codes.get(item_supplier), item_supplier, item.brand or None, item.vat,
        ])
        _apply_flags(ws, r, row.flags)

    for col, fmt in NUMBER_FORMATS.items():
        for cell in ws[col][1:]:
            cell.number_format = fmt
    for i, width in enumerate(WIDTHS):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = width
    ws.freeze_panes = "A2"
    ws.sheet_view.topLeftCell = "A1"

    save_atomic(path, wb.save)


def _apply_flags(ws, r: int, flags: list[Flag]) -> None:
    """Colour the flagged cells and explain each colour in a cell comment (hover to read)."""
    by_cell: dict[str, list[Flag]] = {}
    for flag in flags:
        letters = FLAG_COLUMNS.get(flag.column) or [openpyxl.utils.get_column_letter(c) for c in range(1, 21)]
        for letter in letters:
            by_cell.setdefault(letter, []).append(flag)
    for letter, cell_flags in by_cell.items():
        worst = max(cell_flags, key=lambda f: SEVERITY[f.level])
        cell = ws[f"{letter}{r}"]
        cell.fill = FILLS[worst.level]
    # One comment per row, on the description, listing every reason.
    if flags:
        text = "\n".join(f"[{f.level.upper()}] {f.message}" for f in sorted(flags, key=lambda f: -SEVERITY[f.level]))
        comment = Comment(text, "Invoice tool")
        comment.width, comment.height = 420, 30 + 18 * len(flags)
        ws[f"E{r}"].comment = comment
