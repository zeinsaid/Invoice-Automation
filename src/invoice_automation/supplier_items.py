"""What the tool has learned about each supplier's products (D-018 point 2, D-006 step 1).

One row per (supplier, key). The key is either the barcode the supplier prints, or the supplier's
wording for the product (normalised). A row says which of our items it is, how many selling
pieces one invoice unit holds (pack factor, D-005/D-020), and the supplier's last cost.

Stored as an Excel file so Zein can open, read and correct it. All key/code columns are text,
so barcodes keep their leading zeros. To confirm a match by hand: set Status to 'trusted' (and
fill in Pack factor if it's empty), save, and re-run.
"""

from dataclasses import dataclass, fields
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from .catalogue import name_key
from .files import read_rows, save_atomic
from .layout import ean_is_valid, normalise_barcode

TABLE_PATH = Path("data/supplier_items.xlsx")
LEGACY_PATH = Path("data/reference/legacy_supplier_lists.xlsx")

TRUSTED, PROVISIONAL = "trusted", "provisional"

# Old per-supplier Excel name -> PowerSoft supplier name (for seeding from the legacy lists).
LEGACY_SUPPLIERS = {
    "Athenites": "Chr. Hadjichristofi (Athienitis) Ltd",
    "Agrotis": "Renos Agrotis & Son Ltd",
    "Pambos": "Pambos & Andreas Trading Ltd",
    "Rainbow": "Rainbow Agencies LTD",
    "Fereos": "Phoenicia Fereos Cyprus Ltd",
    "Cassandra": "Cassandra Trading LTD",
    "Vassos": "Vassos Eliades Ltd",
    "Andrikian": "Andrikian Trading Co. LTD",
    "Cosmos": "Cosmos Tobacco Trading LTD",
    "ED Sport": "ED SPORT COMPANY LTD",
    "Argosy": "Argosy Trading Company LTD",
    "Coca-Cola": "Coca-Cola HBC Cyprus",
    "Corina": "Corina Snacks LTD",
    "H&D": "H&D HI-TECH LIGHTING & ENERGY TRADING LTD",
    "Aisian": "Asian Food Market LTD",
    "Jumbo": "Jumbo Trading LTD",
    "Kypropharm": "Kypropharm LTD",
    "OfficeServe": "A&N OFFICESERV BUSINESS SOLUTIONS LTD",
    "Finerico": "FINERICO.LTD",
    "Paradisiotis": "Paradisiotis Ltd",
    "Charalambides": "CHARALAMBIDES CHRISTIS LTD",
    "Aristidou": "Aristidou Ioannis & Sons Water Supply LTD",
    "Party Ice": "Party Ice & Party Time Ice LTD",
}


@dataclass
class SupplierItem:
    supplier: str
    key_type: str  # "barcode" or "name"
    key: str
    supplier_description: str
    item_code: str
    pack_factor: float | None  # selling pieces per invoice unit; None = not known yet
    status: str  # TRUSTED or PROVISIONAL
    source: str  # how it was learned: barcode, name, fuzzy, legacy, manual
    last_cost: float | None  # per selling piece, excl. VAT
    last_invoice: str
    previous_cost: float | None  # the cost before last_invoice (keeps re-runs of one invoice stable)
    last_seen: str


COLUMNS = [f.name for f in fields(SupplierItem)]
HEADER = ["Supplier", "Key type", "Key", "Supplier description", "Item code", "Pack factor",
          "Status", "Source", "Last cost", "Last invoice", "Previous cost", "Last seen"]


class SupplierItems:
    def __init__(self, rows: list[SupplierItem], path: Path = TABLE_PATH):
        self.rows = rows
        self.path = path
        self._index = {(r.supplier, r.key_type, r.key): r for r in rows}

    @classmethod
    def load(cls, path: Path = TABLE_PATH) -> "SupplierItems":
        if not path.exists():
            table = cls(_seed_from_legacy(), path)
            table.save()
            return table
        rows = []
        for values in read_rows(path):
            if not values or not values[0]:
                continue
            d = dict(zip(COLUMNS, values))
            for k in ("pack_factor", "last_cost", "previous_cost"):
                d[k] = float(d[k]) if d[k] not in (None, "") else None
            for k in ("supplier", "key_type", "key", "supplier_description", "item_code", "status",
                      "source", "last_invoice", "last_seen"):
                d[k] = "" if d[k] is None else str(d[k])
            rows.append(SupplierItem(**d))
        return cls(rows, path)

    def get(self, supplier: str, key_type: str, key: str) -> SupplierItem | None:
        return self._index.get((supplier, key_type, key))

    def names_for(self, supplier: str) -> list[SupplierItem]:
        return [r for r in self.rows if r.supplier == supplier and r.key_type == "name"]

    def last_cost(self, supplier: str, item_code: str, invoice: str) -> float | None:
        """This supplier's most recent cost for the item, ignoring the invoice being processed."""
        rows = [r for r in self.rows if r.supplier == supplier and r.item_code == item_code]
        earlier = [(r.last_seen, r.last_cost) for r in rows if r.last_invoice != invoice and r.last_cost is not None]
        if earlier:
            return max(earlier)[1]
        # Only this invoice has set a cost: use what was there before it (re-runs stay stable).
        return next((r.previous_cost for r in rows if r.previous_cost is not None), None)

    def upsert(self, row: SupplierItem) -> SupplierItem:
        """Add a row, or update the existing one with the same (supplier, key type, key)."""
        existing = self.get(row.supplier, row.key_type, row.key)
        if existing is None:
            self.rows.append(row)
            self._index[(row.supplier, row.key_type, row.key)] = row
            return row
        if row.last_invoice and row.last_invoice != existing.last_invoice:
            existing.previous_cost = existing.last_cost
        for k in ("supplier_description", "last_cost", "last_invoice", "last_seen"):
            if getattr(row, k) not in (None, ""):
                setattr(existing, k, getattr(row, k))
        # Never downgrade a trusted row, and never overwrite a factor Zein filled in.
        if existing.status != TRUSTED:
            existing.status = row.status
            existing.item_code = row.item_code
            existing.source = row.source
        if existing.pack_factor is None:
            existing.pack_factor = row.pack_factor
        return existing

    def save(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Supplier items"
        ws.append(HEADER)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="C0C0C0")
        for r in sorted(self.rows, key=lambda r: (r.supplier, r.key_type, r.supplier_description)):
            ws.append([getattr(r, c) for c in COLUMNS])
        for row in ws.iter_rows(min_row=2):
            for i in (2, 4):  # Key, Item code: text, so leading zeros survive
                row[i].number_format = "@"
        for col, width in zip("ABCDEFGHIJKL", (34, 9, 32, 34, 9, 9, 11, 9, 9, 13, 9, 11)):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        save_atomic(self.path, wb.save)  # a power cut mid-save must never cost what was learned


def _seed_from_legacy() -> list[SupplierItem]:
    """First run only: each supplier's own wording for our items, from the old masters (D-018)."""
    rows: dict[tuple, SupplierItem] = {}
    for file_name, code, description, _barcode, cost, _code in read_rows(LEGACY_PATH, first_sheet=True):
        supplier = LEGACY_SUPPLIERS.get(file_name)
        if not supplier or code is None or not description:
            continue
        key = name_key(str(description))
        rows[(supplier, key)] = SupplierItem(
            supplier=supplier, key_type="name", key=key, supplier_description=str(description),
            item_code=str(code), pack_factor=None, status=TRUSTED, source="legacy",
            last_cost=float(cost) if isinstance(cost, (int, float)) else None,
            last_invoice="", previous_cost=None, last_seen=str(date(2026, 10, 14)),
        )
    return list(rows.values())


LEGACY_BARCODE = "legacy barcode"  # source of the rows below


def names_related(a: str, b: str) -> bool:
    """Do two names for one product share any word? Abbreviations count by their first 4 letters
    ('CHOC/' and 'CHOCOLATE'); numbers and short words don't count. A name with no such words passes."""
    def words(name):
        return {w[:4] for w in name_key(name).split() if len(w) > 2 and not any(c.isdigit() for c in w)}
    wa, wb = words(a), words(b)
    return not wa or not wb or bool(wa & wb)


def add_legacy_barcodes(table: SupplierItems, item_names: dict[str, str], powersoft_barcodes: set[str]) -> int:
    """Once (D-028): the old masters also hold each supplier's own barcode for an item, which PowerSoft
    often doesn't have (Coca-Cola prints its can's barcode; PowerSoft has the shelf's). Each one becomes
    an unconfirmed answer, so the line asks "Is it this item?" instead of "not found". `item_names` maps
    PowerSoft item code -> name. Skipped: barcodes PowerSoft already has (those match today, with no
    question), barcodes on two items, items no longer in PowerSoft, misread ones, and rows whose name has
    nothing in common with the item's (the old lists have a few shifted rows: 'NUTELLA JAR 750GR' ->
    'GEEKBAR KIT OFFER'). Returns how many were added; the caller saves the table and runs this once."""
    found: dict[tuple[str, str], set[str]] = {}
    details: dict[tuple[str, str], tuple] = {}
    for file_name, code, description, barcode, cost, _code in read_rows(LEGACY_PATH, first_sheet=True):
        supplier = LEGACY_SUPPLIERS.get(file_name)
        barcode = normalise_barcode(str(barcode or ""))
        if supplier and code is not None and barcode:
            found.setdefault((supplier, barcode), set()).add(str(code))
            details[(supplier, barcode)] = (str(description or ""), cost)
    added = 0
    for (supplier, barcode), codes in found.items():
        code = next(iter(codes))
        description, cost = details[(supplier, barcode)]
        if (len(codes) > 1 or code not in item_names or not ean_is_valid(barcode) or barcode in powersoft_barcodes
                or table.get(supplier, "barcode", barcode) or not names_related(description, item_names[code])):
            continue
        table.upsert(SupplierItem(
            supplier=supplier, key_type="barcode", key=barcode, supplier_description=description, item_code=code,
            pack_factor=None, status=PROVISIONAL, source=LEGACY_BARCODE,
            last_cost=float(cost) if isinstance(cost, (int, float)) else None,
            last_invoice="", previous_cost=None, last_seen=str(date.today())))
        added += 1
    return added
