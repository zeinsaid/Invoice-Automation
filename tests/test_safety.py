"""Data safety (D-023): nothing learned is lost to a crash, a bad export or a missing file."""

import csv
import zipfile

import openpyxl
import pytest

from invoice_automation.backup import backup_daily
from invoice_automation.catalogue import EXPORT_COLUMNS, Catalogue
from invoice_automation.files import read_rows, save_atomic
from invoice_automation.pipeline import Context, move_off_inactive
from invoice_automation.supplier_items import SupplierItem, SupplierItems


def write_export(folder, day, n, columns=EXPORT_COLUMNS):
    with (folder / f"items_export_2026-10-{day:02d}.csv").open("w", encoding="cp1253", newline="") as f:
        w = csv.writer(f)
        w.writerow(columns)
        for i in range(1, n + 1):
            row = {"Item Code": str(i), "Item Description": f"ITEM {i}", "Barcode": "", "VAT": "19%",
                   "Cost": "1.00", "Price 1": "2.00", "Supplier": "S", "Brand": "", "Category": "C", "Storable": "Yes"}
            w.writerow([row.get(c, "") for c in columns])


@pytest.fixture
def powersoft(tmp_path):
    for name in ("categories.xlsx", "suppliers.xlsx"):
        wb = openpyxl.Workbook()
        wb.active.append(["header"])
        wb.save(tmp_path / name)
    return tmp_path


def test_newest_good_export_is_used(powersoft):
    write_export(powersoft, 1, 100)
    write_export(powersoft, 2, 102)
    cat = Catalogue.load(powersoft)
    assert (cat.export_name, len(cat.items), cat.warnings) == ("items_export_2026-10-02.csv", 102, [])


def test_cut_off_export_falls_back_to_the_previous_one(powersoft):
    write_export(powersoft, 1, 100)
    write_export(powersoft, 2, 40)
    cat = Catalogue.load(powersoft)
    assert (cat.export_name, len(cat.items)) == ("items_export_2026-10-01.csv", 100)
    assert "cut off" in cat.warnings[0]


def test_export_missing_a_column_falls_back_to_the_previous_one(powersoft):
    write_export(powersoft, 1, 100)
    write_export(powersoft, 2, 100, [c for c in EXPORT_COLUMNS if c != "Storable"])
    cat = Catalogue.load(powersoft)
    assert cat.export_name == "items_export_2026-10-01.csv" and "Storable" in cat.warnings[0]


def test_inactive_items_are_never_matched_and_their_codes_stay_taken(powersoft):
    # D-031: an export with the Active column. Cola 1 was swapped for the item coded with its barcode; 3 is just inactive.
    with (powersoft / "items_export_2026-10-01.csv").open("w", encoding="cp1253", newline="") as f:
        w = csv.writer(f)
        w.writerow(EXPORT_COLUMNS[:3] + ["Active"] + EXPORT_COLUMNS[3:])
        w.writerow(["1", "COLA", "17", "No", "19%", "1.00", "2.00", "S", "", "C", "Yes"])
        w.writerow(["5449000000439", "COLA", "5449000000439", "Yes", "19%", "1.00", "2.00", "S", "", "C", "Yes"])
        w.writerow(["3", "POD", "3", "No", "19%", "1.00", "2.00", "S", "", "C", "Yes"])
    cat = Catalogue.load(powersoft)
    assert list(cat.items) == ["5449000000439"] and set(cat.inactive) == {"1", "3"}
    assert cat.highest_numeric_code() == 3  # the next new item is 4, not 5449000000440
    assert cat.by_barcode == {"5449000000439": ["5449000000439"]}
    assert (cat.active_twin("1"), cat.active_twin("3")) == ("5449000000439", None)  # same name: the same product
    rows = [SupplierItem("S", "barcode", key, "", code, 2, "trusted", "staff", None, "", None, "")
            for key, code in (("111", "1"), ("333", "3"))]
    assert move_off_inactive(SupplierItems(rows), cat) == 1
    assert [r.item_code for r in rows] == ["5449000000439", "3"] and "moved" in cat.warnings[0]


def test_failed_save_keeps_the_old_file(tmp_path):
    path = tmp_path / "learned.txt"
    path.write_text("months of answers")

    def crash(tmp):
        tmp.write_text("half")
        raise OSError("power cut")

    with pytest.raises(OSError):
        save_atomic(path, crash)
    assert path.read_text() == "months of answers" and list(tmp_path.iterdir()) == [path]


def test_daily_backup_once_a_day(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "data/state").mkdir(parents=True)
    (tmp_path / "data/supplier_items.xlsx").write_text("table")
    (tmp_path / "data/state/queue.json").write_text("[]")
    first = backup_daily()
    assert sorted(zipfile.ZipFile(first).namelist()) == ["state/queue.json", "supplier_items.xlsx"]
    assert backup_daily() is None  # already done today


def test_missing_learned_table_stops_instead_of_starting_over(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "data/state").mkdir(parents=True)
    (tmp_path / "data/state/processed_invoices.csv").write_text("supplier,invoice\n")
    with pytest.raises(FileNotFoundError, match="backups"):
        Context.load()


def test_itemslist_is_taken_in_once_as_a_dated_copy(powersoft):
    write_export(powersoft, 1, 100)
    write_export(powersoft, 2, 101)
    (powersoft / "items_export_2026-10-02.csv").rename(powersoft / "itemsList.csv")  # PowerSoft's own name
    first = Catalogue.load(powersoft)
    assert len(first.items) == 101 and first.export_name.startswith("items_export_") and first.export_name != \
        "items_export_2026-10-01.csv"
    copies = sorted(p.name for p in powersoft.glob("items_export_*.csv"))
    Catalogue.load(powersoft)  # the same itemsList again: no new copy
    assert sorted(p.name for p in powersoft.glob("items_export_*.csv")) == copies


def test_marked_files_leave_ready_to_upload_and_the_rest_stay(tmp_path, monkeypatch):
    from invoice_automation import pipeline
    monkeypatch.chdir(tmp_path)
    ready = tmp_path / "output/Ready to upload"
    ready.mkdir(parents=True)
    for name in ("A 1.xlsx", "B 2.xlsx"):
        (ready / name).write_text("x")
    pipeline._save_to_upload({"A 1.xlsx": {"items_import": "1 new item", "made": "2026-10-18"},
                              "B 2.xlsx": {"items_import": "", "made": "2026-10-19"}})
    assert [w["file"] for w in pipeline.waiting_uploads()] == ["A 1.xlsx", "B 2.xlsx"]  # oldest first
    pipeline.mark_uploaded(["A 1.xlsx"])
    assert [w["file"] for w in pipeline.waiting_uploads()] == ["B 2.xlsx"]
    assert len(list((tmp_path / "output/Uploaded").rglob("A 1.xlsx"))) == 1


def test_mark_uploaded_with_the_file_open_leaves_no_copy(tmp_path, monkeypatch):
    # Excel locks an open file: before, the move copied it to Uploaded and then failed to delete it.
    from pathlib import Path

    from invoice_automation import pipeline
    monkeypatch.chdir(tmp_path)
    ready = tmp_path / "output/Ready to upload"
    ready.mkdir(parents=True)
    (ready / "A 1.xlsx").write_text("x")

    def locked(self, target):
        raise PermissionError("in use")
    monkeypatch.setattr(Path, "rename", locked)
    with pytest.raises(PermissionError, match="open in Excel"):
        pipeline.mark_uploaded(["A 1.xlsx"])
    assert (ready / "A 1.xlsx").exists() and not list((tmp_path / "output/Uploaded").rglob("*.xlsx"))


def test_old_supplier_barcodes_are_taken_in_unconfirmed(tmp_path, monkeypatch):
    # D-028: a supplier's own barcode becomes an unconfirmed answer, unless PowerSoft already has it.
    from invoice_automation import supplier_items as si
    legacy = tmp_path / "legacy.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["file", "code", "description", "barcode", "cost", "supplier code"])
    for row in (["Coca-Cola", 21, "FANTA LEMON CAN 330ML", "5449000317667", 0.6, 1],   # Coca-Cola's own: taken in
                ["Coca-Cola", 33, "SPRITE 1.5 LTR", "5449000004826", 1.6, 1],          # PowerSoft has it: skipped
                ["Coca-Cola", 99, "GONE ITEM", "5449000000439", 1.0, 1],               # not in PowerSoft: skipped
                ["Coca-Cola", 21, "FANTA LEMON", "5449000317668", 0.6, 1]):            # bad check digit: skipped
        wb.active.append(row)
    wb.save(legacy)
    monkeypatch.setattr(si, "LEGACY_PATH", legacy)
    table = si.SupplierItems([], tmp_path / "t.xlsx")
    names = {"21": "FANTA LEMON CAN 330ML", "33": "SPRITE 1.5 LTR"}
    assert si.add_legacy_barcodes(table, names, {"5449000004826"}) == 1
    [row] = table.rows
    assert (row.key, row.item_code, row.status) == ("5449000317667", "21", si.PROVISIONAL)


def test_names_related():
    from invoice_automation.supplier_items import names_related
    assert names_related("FITNESS CHOCO/BARS 6X23,5G", "Fitness Bar Chocolate 23.5Gr")
    assert not names_related("NUTELLA JAR 750GR", "GEEKBAR KIT OFFER")  # a shifted row in the old lists


def test_a_table_just_read_can_be_saved_at_once(tmp_path):
    # A read-only workbook kept the learned table open, so saving it straight after loading failed.
    path = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["header"])
    wb.active.append(["row"])
    wb.save(path)
    assert read_rows(path) == [("row",)]
    save_atomic(path, wb.save)  # raises PermissionError ("open in Excel") on Windows if the file were still open
